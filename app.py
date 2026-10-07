"""
app.py  –  SolarAI Ban Pong  |  Flask Application Entry Point
Decision Support System for Solar Energy Analysis in Ban Pong, Ratchaburi
"""

import os
import json
import datetime
import functools
from flask import (
    Flask, render_template, request, jsonify,
    send_file, session, redirect, url_for, flash
)

import predict as pr
import financial_engine as fe
from config import Config

app = Flask(__name__)
app.config.from_object(Config)
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0

BASE_DIR      = os.path.dirname(__file__)
HISTORY_FILE  = os.path.join(BASE_DIR, 'predictions.json')
SETTINGS_FILE = os.path.join(BASE_DIR, 'settings.json')

DEFAULT_SETTINGS = {
    'electricity_rate': 4.50,
    'pea_buyback_rate': 2.20,
    'daytime_ratio': 0.65,
    'co2_factor': 0.4985,
    'typical_monthly_usage': 400,
    'study_area': {
        'district': 'อำเภอบ้านโป่ง',
        'province': 'จังหวัดราชบุรี',
        'lat': 13.8264,
        'lng': 99.8780,
        'zoom': 13
    },
    'solar_packages': [
        {
            'id': 'pkg-3kw',
            'kw': 3,
            'name': 'Solar Rooftop 3 kW (Eco Lite)',
            'cost': 95000,
            'inverter': 'Huawei SUN2000-3KTL / Sungrow SG3.0RS (1 Phase)',
            'panels': 'Tier-1 Mono Half-Cell 585W (5-6 แผง)',
            'warranty_inverter': '10 ปี',
            'warranty_panel': '25 ปี (Linear Performance)',
            'recommended_bill_range': '1,500 – 3,500 บาท/เดือน',
            'description': 'เหมาะสำหรับบ้านพักอาศัยขนาดเล็ก แอร์ 1-2 ตัว ใช้งานกลางวัน',
            'is_active': True
        },
        {
            'id': 'pkg-5kw',
            'kw': 5,
            'name': 'Solar Rooftop 5 kW (Standard Popular)',
            'cost': 150000,
            'inverter': 'Huawei SUN2000-5KTL / Sungrow SG5.0RS (1/3 Phase)',
            'panels': 'Tier-1 Mono Half-Cell 585W (9 แผง)',
            'warranty_inverter': '10 ปี',
            'warranty_panel': '25 ปี (Linear Performance)',
            'recommended_bill_range': '3,800 – 7,500 บาท/เดือน',
            'description': 'รุ่นยอดนิยมสำหรับบ้านเดี่ยว แอร์ 2-4 ตัว ประหยัดสูงสุด',
            'is_active': True
        },
        {
            'id': 'pkg-10kw',
            'kw': 10,
            'name': 'Solar Rooftop 10 kW (Premium Pro)',
            'cost': 250000,
            'inverter': 'Huawei SUN2000-10KTL / Sungrow SG10RT (3 Phase)',
            'panels': 'Tier-1 Mono Half-Cell 585W (18 แผง)',
            'warranty_inverter': '10 ปี',
            'warranty_panel': '25 ปี (Linear Performance)',
            'recommended_bill_range': '7,500 – 13,000 บาท/เดือน',
            'description': 'สำหรับบ้านหลังใหญ่ โฮมออฟฟิศ แอร์ 4-6 ตัว หรือมี EV Charger',
            'is_active': True
        },
        {
            'id': 'pkg-15kw',
            'kw': 15,
            'name': 'Solar Rooftop 15 kW (Commercial Max)',
            'cost': 350000,
            'inverter': 'Huawei SUN2000-15KTL-M2 (3 Phase)',
            'panels': 'Tier-1 Mono Half-Cell 585W (26 แผง)',
            'warranty_inverter': '10 ปี',
            'warranty_panel': '25 ปี (Linear Performance)',
            'recommended_bill_range': '13,000 – 18,000 บาท/เดือน',
            'description': 'สำหรับธุรกิจขนาดกลาง สำนักงาน คลินิก ร้านกาแฟ ใช้งานไฟฟ้าต่อเนื่อง',
            'is_active': True
        },
        {
            'id': 'pkg-20kw',
            'kw': 20,
            'name': 'Solar Rooftop 20 kW (Enterprise Ultra)',
            'cost': 440000,
            'inverter': 'Huawei SUN2000-20KTL-M2 / Sungrow SG20RT (3 Phase)',
            'panels': 'Tier-1 Mono Half-Cell 585W (35 แผง)',
            'warranty_inverter': '10 ปี',
            'warranty_panel': '25 ปี (Linear Performance)',
            'recommended_bill_range': '18,000 – 30,000+ บาท/เดือน',
            'description': 'สำหรับอาคารพาณิชย์ โรงงานขนาดย่อม โรงแรม หอพัก ใช้ไฟช่วงกลางวันสูง',
            'is_active': True
        }
    ],
    'system_costs': {'3': 95000, '5': 150000, '10': 250000, '15': 350000, '20': 440000},
    'admin_user': 'admin',
    'admin_pass': 'solarai2026'
}


# ── File I/O helpers ─────────────────────────────────────────────────────────

def load_json(path, default):
    if os.path.exists(path):
        try:
            with open(path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return default


def save_json(path, data):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def get_settings():
    saved = load_json(SETTINGS_FILE, {})
    return {**DEFAULT_SETTINGS, **saved}


def get_history():
    return load_json(HISTORY_FILE, [])


def save_history(history):
    save_json(HISTORY_FILE, history)


# ── Context Processor ─────────────────────────────────────────────────────────

@app.context_processor
def inject_global_context():
    return {
        'is_admin': session.get('is_admin', False),
        'admin_username': session.get('admin_username', 'Administrator'),
        'current_year': datetime.datetime.now().year,
        'app_settings': get_settings()
    }


# ── Auth Decorator ───────────────────────────────────────────────────────────

def admin_required(f):
    @functools.wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('is_admin'):
            flash('กรุณาเข้าสู่ระบบผู้ดูแลระบบก่อนเข้าใช้งานส่วนนี้', 'warning')
            return redirect(url_for('admin_login', next=request.url))
        return f(*args, **kwargs)
    return decorated_function


# ── Public User Routes ───────────────────────────────────────────────────────

@app.route('/')
def dashboard():
    """User Dashboard / Overview."""
    return render_template('index.html', active='dashboard')


@app.route('/map')
def gis_map():
    """GIS Solar Mapping & Analysis for Ban Pong."""
    return render_template('map.html', active='map')


@app.route('/analysis/<analysis_id>')
def analysis_detail(analysis_id):
    """Detailed solar feasibility and financial breakdown."""
    history = get_history()
    record  = next((r for r in history if str(r.get('id')) == str(analysis_id)), None)
    if record is None:
        return redirect(url_for('gis_map'))
    return render_template('analysis.html', active='map',
                           data=record,
                           data_json=json.dumps(record, ensure_ascii=False))


@app.route('/history')
def history_page():
    """User historical calculations log."""
    return render_template('history.html', active='history')


@app.route('/report')
def report_page():
    """Solar Feasibility Reports Generator."""
    history = get_history()
    return render_template('report.html', active='report',
                           history=list(reversed(history))[:25])


@app.route('/report/document/<analysis_id>')
@app.route('/report/print/<analysis_id>')
def report_document(analysis_id):
    """Official printable / PDF-ready Solar Feasibility Report."""
    history = get_history()
    record  = next((r for r in history if str(r.get('id')) == str(analysis_id)), None)
    if record is None:
        return redirect(url_for('report_page'))
    return render_template('document_report.html',
                           data=record,
                           data_json=json.dumps(record, ensure_ascii=False))


# ── Admin Routes ─────────────────────────────────────────────────────────────

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    """Admin Login Page."""
    if session.get('is_admin'):
        return redirect(url_for('admin_dashboard'))

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        cfg = get_settings()

        expected_user = cfg.get('admin_user', 'admin')
        expected_pass = cfg.get('admin_pass', 'solarai2026')

        if username == expected_user and password == expected_pass:
            session['is_admin'] = True
            session['admin_username'] = username
            session.permanent = True
            flash('ยินดีต้อนรับเข้าสู่ระบบผู้ดูแลระบบ SolarAI', 'success')
            next_url = request.args.get('next') or url_for('admin_dashboard')
            return redirect(next_url)
        else:
            flash('ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง กรุณาลองใหม่อีกครั้ง', 'error')

    return render_template('admin_login.html', active='admin_login')


@app.route('/admin/logout')
def admin_logout():
    """Admin Logout."""
    session.pop('is_admin', None)
    session.pop('admin_username', None)
    flash('ออกจากระบบผู้ดูแลระบบเรียบร้อยแล้ว', 'info')
    return redirect(url_for('dashboard'))


@app.route('/admin/dashboard')
@admin_required
def admin_dashboard():
    """Admin Executive Dashboard with System Health, Stats & Analytics."""
    history = get_history()
    settings = get_settings()
    model_info = pr.get_model_info()

    # Aggregate stats
    total_runs = len(history)
    total_energy_kwh = sum((r.get('annual_kwh') or 0) for r in history)
    total_energy_mwh = round(total_energy_kwh / 1000.0, 2)
    total_savings_thb = sum((r.get('annual_save') or 0) for r in history)
    total_co2_tons = round(sum((r.get('co2_saved_kg') or 0) for r in history) / 1000.0, 2)
    
    avg_roi = round(sum((r.get('roi') or 0) for r in history) / total_runs, 1) if total_runs > 0 else 0.0
    avg_breakeven = round(sum((r.get('breakeven') or 0) for r in history) / total_runs, 1) if total_runs > 0 else 0.0

    # System distribution
    package_counts = {}
    for r in history:
        kw = str(r.get('recommended', {}).get('kw', 5))
        package_counts[kw] = package_counts.get(kw, 0) + 1

    return render_template('admin_dashboard.html',
                           active='admin_dashboard',
                           history=list(reversed(history)),
                           settings=settings,
                           model_info=model_info,
                           total_runs=total_runs,
                           total_energy_mwh=total_energy_mwh,
                           total_savings_thb=total_savings_thb,
                           total_co2_tons=total_co2_tons,
                           avg_roi=avg_roi,
                           avg_breakeven=avg_breakeven,
                           package_counts=package_counts)


@app.route('/admin/packages')
@admin_required
def admin_packages():
    """Admin CRUD for Solar Packages & Hardware Specs."""
    settings = get_settings()
    packages = settings.get('solar_packages', [])
    return render_template('admin_packages.html',
                           active='admin_packages',
                           packages=packages,
                           settings=settings)


@app.route('/admin/audit')
@admin_required
def admin_audit():
    """Admin Audit Trail & Inspection of Predictions."""
    history = get_history()
    return render_template('admin_audit.html',
                           active='admin_audit',
                           history=list(reversed(history)))


@app.route('/admin/settings')
@app.route('/settings')
@admin_required
def settings_page():
    """Settings page (strictly restricted to Admin)."""
    return render_template('settings.html', active='settings',
                           settings=get_settings())


# ── Public API routes ────────────────────────────────────────────────────────

@app.route('/api/boundary')
def api_boundary():
    """Serve the Ban Pong district boundary GeoJSON."""
    path = os.path.join(BASE_DIR, 'static', 'banpong.geojson')
    return send_file(path, mimetype='application/json')


@app.route('/api/analyze', methods=['POST'])
def api_analyze():
    """Run ML Prediction & Financial Engineering."""
    body         = request.get_json(silent=True) or {}
    lat          = float(body.get('lat', 13.8264))
    lng          = float(body.get('lng', 99.8780))
    address      = body.get('address', f'{lat:.5f}, {lng:.5f}')
    monthly_bill = float(body.get('monthly_bill', 5000))
    building_type = body.get('building_type', 'บ้านพักอาศัย')
    roof_type     = body.get('roof_type', 'กระเบื้องซีแพค')

    cfg          = get_settings()
    costs        = {int(k): int(v) for k, v in cfg.get('system_costs', {}).items()}
    elec_rate    = float(cfg.get('electricity_rate', 4.50))
    buyback_rate = float(cfg.get('pea_buyback_rate', 2.20))
    # Prefer user-provided daytime_ratio, fallback to admin settings
    daytime_ratio = float(body.get('daytime_ratio', cfg.get('daytime_ratio', 0.65)))

    try:
        # 1. AI prediction (production values for Ban Pong)
        raw = pr.analyze_location(lat, lng, electricity_rate=elec_rate, system_costs=costs)

        # Inject dynamic rates so financial_engine uses admin-configured values
        raw['pea_buyback_rate'] = buyback_rate
        raw['daytime_ratio']    = daytime_ratio

        # 2. Financial engine – calculates payback, ROI, savings
        result = fe.calculate_financials(raw, monthly_bill)

        rec_id = int(datetime.datetime.now().timestamp() * 1000)
        record = {
            'id': rec_id,
            'address': address,
            'building_type': building_type,
            'roof_type': roof_type,
            'daytime_ratio': daytime_ratio,
            **result
        }

        history = get_history()
        history.append(record)
        save_history(history)

        return jsonify({'status': 'ok', **record})
    except Exception as exc:
        app.logger.exception('Analysis error')
        return jsonify({'status': 'error', 'message': str(exc)}), 500


@app.route('/api/analysis/<analysis_id>')
def api_get_analysis(analysis_id):
    history = get_history()
    record  = next((r for r in history if str(r.get('id')) == str(analysis_id)), None)
    if record:
        return jsonify(record)
    return jsonify({'error': 'Not found'}), 404


@app.route('/api/stats')
def api_stats():
    history = get_history()
    if not history:
        return jsonify({'count': 0, 'avg_daily': 0, 'avg_annual': 0,
                        'avg_save': 0, 'avg_roi': 0, 'avg_breakeven': 0})
    n = len(history)

    def avg(key):
        vals = [(r.get(key) or 0) for r in history]
        return round(sum(vals) / n, 2) if n > 0 else 0

    return jsonify({
        'count':         n,
        'avg_daily':     avg('current_daily'),
        'avg_annual':    avg('annual_kwh'),
        'avg_save':      avg('annual_save'),
        'avg_roi':       avg('roi'),
        'avg_breakeven': avg('breakeven'),
    })


@app.route('/api/model-info')
def api_model_info():
    try:
        return jsonify(pr.get_model_info())
    except Exception as exc:
        return jsonify({'error': str(exc)}), 500


@app.route('/api/history', methods=['GET'])
def api_history():
    return jsonify(list(reversed(get_history())))


@app.route('/api/history/<analysis_id>', methods=['DELETE'])
def api_delete(analysis_id):
    history = get_history()
    new     = [r for r in history if str(r.get('id')) != str(analysis_id)]
    if len(new) < len(history):
        save_history(new)
        return jsonify({'status': 'ok'})
    return jsonify({'error': 'Not found'}), 404


@app.route('/api/settings', methods=['GET', 'POST'])
def api_settings():
    if request.method == 'POST':
        if not session.get('is_admin'):
            return jsonify({'error': 'จำเป็นต้องเข้าสู่ระบบผู้ดูแลระบบ (Admin) เพื่อแก้ไขการตั้งค่า'}), 403
        body    = request.get_json(silent=True) or {}
        current = get_settings()
        current.update(body)
        save_json(SETTINGS_FILE, current)
        return jsonify({'status': 'ok', **current})
    return jsonify(get_settings())


# ── Admin API routes ─────────────────────────────────────────────────────────

@app.route('/api/admin/diagnostics')
@admin_required
def api_admin_diagnostics():
    """System health & ML Pipeline diagnostics check."""
    model_loaded = False
    model_r2 = 0.0
    try:
        m = pr.load_model()
        model_loaded = m is not None
        model_info = pr.get_model_info()
        model_r2 = model_info.get('r2_score', 0.94)
    except Exception as e:
        model_loaded = False
        model_info = {'error': str(e)}

    # File integrity checks
    db_predictions = os.path.exists(HISTORY_FILE)
    db_settings = os.path.exists(SETTINGS_FILE)
    geojson_bound = os.path.exists(os.path.join(BASE_DIR, 'static', 'banpong.geojson'))
    
    return jsonify({
        'status': 'healthy' if model_loaded and db_predictions and db_settings else 'warning',
        'model_status': 'online' if model_loaded else 'offline',
        'model_r2': model_r2,
        'features_count': len(model_info.get('features', [])),
        'database_records': len(get_history()),
        'banpong_geojson': geojson_bound,
        'settings_valid': db_settings,
        'server_time': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'auth_user': session.get('admin_username', 'admin')
    })


@app.route('/api/admin/simulate', methods=['POST'])
@admin_required
def api_admin_simulate():
    """Admin live calculation simulation test bench."""
    body = request.get_json(silent=True) or {}
    lat = float(body.get('lat', 13.8264))
    lng = float(body.get('lng', 99.8780))
    monthly_bill = float(body.get('monthly_bill', 5000))
    
    cfg = get_settings()
    costs = {int(k): int(v) for k, v in cfg.get('system_costs', {}).items()}
    elec_rate = float(cfg.get('electricity_rate', 4.50))
    
    raw = pr.analyze_location(lat, lng, electricity_rate=elec_rate, system_costs=costs)
    result = fe.calculate_financials(raw, monthly_bill)
    
    return jsonify({
        'status': 'ok',
        'simulation': result,
        'inputs': {
            'lat': lat,
            'lng': lng,
            'monthly_bill': monthly_bill,
            'electricity_rate': elec_rate,
            'costs': costs
        }
    })

@app.route('/api/admin/packages', methods=['GET', 'POST'])
@admin_required
def api_admin_packages():
    """Get all packages or Create a new package."""
    cfg = get_settings()
    packages = cfg.get('solar_packages', [])

    if request.method == 'POST':
        data = request.get_json(silent=True) or {}
        kw = float(data.get('kw', 5))
        cost = int(data.get('cost', 150000))
        name = data.get('name', f'Solar Package {kw} kW')
        
        new_pkg = {
            'id': f'pkg-{int(kw) if kw.is_integer() else str(kw).replace(".", "_")}kw-{int(datetime.datetime.now().timestamp())}',
            'kw': int(kw) if kw.is_integer() else kw,
            'name': name,
            'cost': cost,
            'inverter': data.get('inverter', 'Huawei / Sungrow Inverter'),
            'panels': data.get('panels', 'Tier-1 Mono Half-Cell 585W'),
            'warranty_inverter': data.get('warranty_inverter', '10 ปี'),
            'warranty_panel': data.get('warranty_panel', '25 ปี'),
            'recommended_bill_range': data.get('recommended_bill_range', 'ตามการประเมิน'),
            'description': data.get('description', ''),
            'is_active': bool(data.get('is_active', True))
        }

        packages.append(new_pkg)
        # update system_costs map
        cfg['solar_packages'] = packages
        sys_costs = cfg.get('system_costs', {})
        sys_costs[str(int(kw) if kw.is_integer() else kw)] = cost
        cfg['system_costs'] = sys_costs

        save_json(SETTINGS_FILE, cfg)
        return jsonify({'status': 'ok', 'package': new_pkg})

    return jsonify(packages)


@app.route('/api/admin/packages/<pkg_id>', methods=['PUT', 'DELETE'])
@admin_required
def api_admin_package_detail(pkg_id):
    """Update or Delete a solar package."""
    cfg = get_settings()
    packages = cfg.get('solar_packages', [])

    idx = next((i for i, p in enumerate(packages) if p.get('id') == pkg_id), -1)
    if idx == -1:
        return jsonify({'error': 'Package not found'}), 404

    if request.method == 'DELETE':
        deleted = packages.pop(idx)
        # update system_costs
        kw_str = str(deleted.get('kw'))
        if kw_str in cfg.get('system_costs', {}):
            cfg['system_costs'].pop(kw_str, None)
        cfg['solar_packages'] = packages
        save_json(SETTINGS_FILE, cfg)
        return jsonify({'status': 'ok', 'deleted_id': pkg_id})

    if request.method == 'PUT':
        data = request.get_json(silent=True) or {}
        pkg = packages[idx]
        pkg['name'] = data.get('name', pkg['name'])
        pkg['cost'] = int(data.get('cost', pkg['cost']))
        pkg['kw'] = float(data.get('kw', pkg['kw']))
        if pkg['kw'].is_integer():
            pkg['kw'] = int(pkg['kw'])
        pkg['inverter'] = data.get('inverter', pkg.get('inverter', ''))
        pkg['panels'] = data.get('panels', pkg.get('panels', ''))
        pkg['warranty_inverter'] = data.get('warranty_inverter', pkg.get('warranty_inverter', ''))
        pkg['warranty_panel'] = data.get('warranty_panel', pkg.get('warranty_panel', ''))
        pkg['recommended_bill_range'] = data.get('recommended_bill_range', pkg.get('recommended_bill_range', ''))
        pkg['description'] = data.get('description', pkg.get('description', ''))
        pkg['is_active'] = bool(data.get('is_active', True))

        # update system_costs map
        cfg['system_costs'][str(pkg['kw'])] = pkg['cost']
        cfg['solar_packages'] = packages
        save_json(SETTINGS_FILE, cfg)
        return jsonify({'status': 'ok', 'package': pkg})


@app.route('/api/admin/audit/clear', methods=['POST'])
@admin_required
def api_admin_clear_audit():
    """Clear all historical calculation audit records."""
    save_history([])
    return jsonify({'status': 'ok', 'message': 'ล้างประวัติการวิเคราะห์ทั้งหมดเรียบร้อยแล้ว'})


# ── Excel Export ─────────────────────────────────────────────────────────────

@app.route('/api/export/excel/<analysis_id>')
def api_export_excel(analysis_id):
    history = get_history()

    # Support 'all' to export all historical records into one workbook
    if analysis_id == 'all':
        return export_all_excel(history)

    record = next((r for r in history if str(r.get('id')) == str(analysis_id)), None)
    if not record:
        return jsonify({'error': 'Not found'}), 404

    try:
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import BarChart, Reference

        wb = openpyxl.Workbook()

        # Styles
        font_title = Font(name='Calibri', size=16, bold=True, color='0F172A')
        font_sub   = Font(name='Calibri', size=11, italic=True, color='475569')
        font_hdr   = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        font_bold  = Font(name='Calibri', size=11, bold=True, color='0F172A')
        font_norm  = Font(name='Calibri', size=11, color='1E293B')

        fill_hdr   = PatternFill(start_color='059669', end_color='059669', fill_type='solid') # Emerald 600
        fill_rec   = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid') # Emerald 100
        fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

        thin = Side(border_style='thin', color='CBD5E1')
        border_box = Border(left=thin, right=thin, top=thin, bottom=thin)

        align_left   = Alignment(horizontal='left', vertical='center')
        align_right  = Alignment(horizontal='right', vertical='center')
        align_center = Alignment(horizontal='center', vertical='center')

        # ── Sheet 1: สรุปผล ───────────────────────────────────────────────────
        ws1 = wb.active
        ws1.title = 'สรุปผลการวิเคราะห์'

        ws1['A1'] = 'SolarAI Ban Pong – รายงานการวิเคราะห์พลังงานแสงอาทิตย์'
        ws1['A1'].font = font_title
        ws1['A2'] = f"สถานที่: {record.get('address', '')}  |  วันที่: {record.get('analyzed_date', '')} {record.get('analyzed_time', '')}"
        ws1['A2'].font = font_sub

        summary_rows = [
            ('รายการ',                              'ค่าข้อมูล'),
            ('สถานที่วิเคราะห์',                      record.get('address', '')),
            ('พิกัด GPS',                          f"{record.get('lat', 0)}°N, {record.get('lng', 0)}°E"),
            ('วันที่และเวลาที่วิเคราะห์',             f"{record.get('analyzed_date', '')} {record.get('analyzed_time', '')}"),
            ('ระดับความเหมาะสมของพื้นที่',            f"{record.get('suitability', 0)} / 5 ดาว ({record.get('suitability_label', '')})"),
            ('ความมั่นใจโมเดล AI',                   f"{record.get('confidence', 0)}%"),
            ('ระบบโซลาร์เซลล์ที่แนะนำ',               f"{record.get('recommended', {}).get('kw', 0)} kW"),
            ('ราคาระบบที่แนะนำ (บาท)',              record.get('recommended', {}).get('cost', 0)),
            ('ประมาณการผลิตไฟฟ้า / วัน (kWh)',        record.get('current_daily', 0)),
            ('ประมาณการผลิตไฟฟ้า / ปี (kWh)',         record.get('annual_kwh', 0)),
            ('ประหยัดค่าไฟฟ้า / เดือน (บาท)',          record.get('monthly_save', 0)),
            ('ประหยัดค่าไฟฟ้า / ปี (บาท)',             record.get('annual_save', 0)),
            ('ผลตอบแทนการลงทุนต่อปี (ROI)',           f"{record.get('roi', 0)}%"),
            ('ระยะเวลาคืนทุน (ปี)',                  f"{record.get('breakeven', 0)} ปี"),
            ('ประหยัดสะสม 25 ปี (บาท)',               record.get('lifetime_savings_25y', 0)),
            ('กำไรสุทธิ 25 ปี (บาท)',                 record.get('lifetime_net_profit', 0)),
            ('ปริมาณ CO₂ ที่ลดได้ / ปี (kg)',          record.get('co2_saved_kg', 0)),
            ('เทียบเท่าการปลูกต้นไม้ (ต้น/ปี)',       record.get('trees_equiv', 0)),
        ]

        start_row = 4
        for r_idx, (k, v) in enumerate(summary_rows, start=start_row):
            cell_k = ws1.cell(row=r_idx, column=1, value=k)
            cell_v = ws1.cell(row=r_idx, column=2, value=v)

            if r_idx == start_row:
                cell_k.font, cell_v.font = font_hdr, font_hdr
                cell_k.fill, cell_v.fill = fill_hdr, fill_hdr
                cell_k.alignment, cell_v.alignment = align_left, align_right
            else:
                cell_k.font = font_bold if 'แนะนำ' in str(k) else font_norm
                cell_v.font = font_bold if 'แนะนำ' in str(k) else font_norm
                cell_k.alignment, cell_v.alignment = align_left, align_right
                if r_idx % 2 == 1:
                    cell_k.fill, cell_v.fill = fill_zebra, fill_zebra

            cell_k.border, cell_v.border = border_box, border_box

        # ── Sheet 2: รายเดือน ──────────────────────────────────────────────────
        ws2 = wb.create_sheet(title='พยากรณ์รายเดือน')
        ws2['A1'] = 'การผลิตไฟฟ้าพยากรณ์รายเดือน (12 เดือน)'
        ws2['A1'].font = font_title

        headers_m = ['เดือน', 'ฤดูกาล', 'ผลิต/วัน (kWh)', 'ผลิต/เดือน (kWh)', 'ประหยัด/เดือน (บาท)']
        for c_idx, h in enumerate(headers_m, start=1):
            cell = ws2.cell(row=3, column=c_idx, value=h)
            cell.font, cell.fill, cell.border = font_hdr, fill_hdr, border_box
            cell.alignment = align_center if c_idx == 2 else (align_left if c_idx == 1 else align_right)

        monthly = record.get('monthly_data', [])
        rate = record.get('electricity_rate', 4.50)
        for r_idx, m in enumerate(monthly, start=4):
            m_kwh = m.get('monthly_kwh', 0)
            row_data = [
                m.get('month_name', ''),
                m.get('season', ''),
                m.get('daily_kwh', 0),
                m_kwh,
                round(min(record.get('monthly_bill', 999999), m_kwh * rate))
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws2.cell(row=r_idx, column=c_idx, value=val)
                cell.font, cell.border = font_norm, border_box
                cell.alignment = align_center if c_idx == 2 else (align_left if c_idx == 1 else align_right)
                if r_idx % 2 == 1:
                    cell.fill = fill_zebra

        # Total Row
        tot_row = len(monthly) + 4
        tot_kwh = sum(m.get('monthly_kwh', 0) for m in monthly)
        tot_save = record.get('annual_save', 0)
        tot_vals = ['รวมทั้งปี', '-', round(record.get('current_daily', 0), 2), round(tot_kwh, 1), tot_save]
        for c_idx, val in enumerate(tot_vals, start=1):
            cell = ws2.cell(row=tot_row, column=c_idx, value=val)
            cell.font = font_bold
            cell.fill = fill_rec
            cell.border = border_box
            cell.alignment = align_center if c_idx == 2 else (align_left if c_idx == 1 else align_right)

        # Chart
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "การผลิตไฟฟ้าพยากรณ์รายเดือน (kWh)"
        chart.y_axis.title = 'kWh'
        chart.x_axis.title = 'เดือน'
        chart.width = 14
        chart.height = 7.5

        data_ref = Reference(ws2, min_col=4, min_row=3, max_row=len(monthly)+3, max_col=4)
        cats_ref = Reference(ws2, min_col=1, min_row=4, max_row=len(monthly)+3)
        
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.legend = None

        ws2.add_chart(chart, "G4")

        # ── Sheet 3: เปรียบเทียบขนาดระบบ ───────────────────────────────────────
        ws3 = wb.create_sheet(title='เปรียบเทียบขนาดระบบ')
        ws3['A1'] = 'เปรียบเทียบขนาดระบบโซลาร์เซลล์ (3, 5, 10, 15, 20 kW)'
        ws3['A1'].font = font_title

        headers_s = ['ขนาดระบบ (kW)', 'ผลิต/ปี (kWh)', 'ราคาติดตั้ง (บาท)', 'ประหยัด/ปี (บาท)', 'ROI (%)', 'คืนทุน (ปี)', 'สถานะ']
        for c_idx, h in enumerate(headers_s, start=1):
            cell = ws3.cell(row=3, column=c_idx, value=h)
            cell.font, cell.fill, cell.border = font_hdr, fill_hdr, border_box
            cell.alignment = align_center if c_idx in (1, 7) else align_right

        systems = record.get('systems', [])
        for r_idx, s in enumerate(systems, start=4):
            is_rec = s.get('recommended', False)
            status_text = '⭐ แนะนำ' if is_rec else 'ทางเลือก'
            row_data = [
                f"{s.get('kw', 0)} kW",
                s.get('annual_kwh', 0),
                s.get('cost', 0),
                s.get('annual_save', 0),
                s.get('roi', 0),
                s.get('breakeven', 0),
                status_text
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws3.cell(row=r_idx, column=c_idx, value=val)
                cell.font = font_bold if is_rec else font_norm
                cell.border = border_box
                cell.alignment = align_center if c_idx in (1, 7) else align_right
                if is_rec:
                    cell.fill = fill_rec

        # Auto-adjust column widths
        for sheet in [ws1, ws2, ws3]:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 14)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f'SolarAI_BanPong_{analysis_id}.xlsx'
        resp = send_file(
            buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
        resp.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        return resp
    except Exception as exc:
        app.logger.exception('Excel export error')
        return jsonify({'error': str(exc)}), 500


def export_all_excel(history):
    """Export summary of all history records into one Excel file."""
    try:
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'ประวัติการวิเคราะห์ทั้งหมด'

        font_title = Font(name='Calibri', size=16, bold=True, color='0F172A')
        font_hdr   = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        font_norm  = Font(name='Calibri', size=11, color='1E293B')
        fill_hdr   = PatternFill(start_color='059669', end_color='059669', fill_type='solid')
        fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        thin       = Side(border_style='thin', color='CBD5E1')
        border_box = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws['A1'] = 'SolarAI Ban Pong – รายงานสรุปการวิเคราะห์ทั้งหมด'
        ws['A1'].font = font_title

        headers = ['ID', 'สถานที่', 'วันที่วิเคราะห์', 'ความเหมาะสม', 'ขนาดแนะนำ (kW)', 'ผลิต/วัน (kWh)', 'ผลิต/ปี (kWh)', 'ประหยัด/ปี (บาท)', 'ROI (%)', 'คืนทุน (ปี)']
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=c_idx, value=h)
            cell.font, cell.fill, cell.border = font_hdr, fill_hdr, border_box
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for r_idx, r in enumerate(history, start=4):
            row_data = [
                str(r.get('id', '')),
                r.get('address', ''),
                f"{r.get('analyzed_date', '')} {r.get('analyzed_time', '')}",
                f"{r.get('suitability', 0)}/5 ดาว ({r.get('suitability_label', '')})",
                f"{r.get('recommended', {}).get('kw', 5)} kW",
                r.get('current_daily', 0),
                r.get('annual_kwh', 0),
                r.get('annual_save', 0),
                r.get('roi', 0),
                r.get('breakeven', 0)
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font, cell.border = font_norm, border_box
                cell.alignment = Alignment(horizontal='right' if c_idx >= 6 else 'left', vertical='center')
                if r_idx % 2 == 1:
                    cell.fill = fill_zebra

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f'SolarAI_BanPong_All_Audits_{int(datetime.datetime.now().timestamp())}.xlsx'
        resp = send_file(
            buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
        resp.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
        resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        return resp
    except Exception as exc:
        app.logger.exception('All Excel export error')
        return jsonify({'error': str(exc)}), 500


# ── Entry point ──────────────────────────────────────────────────────────────

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    debug_mode = os.environ.get('FLASK_DEBUG', 'false').lower() in ('1', 'true')
    app.run(host='0.0.0.0', port=port, debug=debug_mode)
