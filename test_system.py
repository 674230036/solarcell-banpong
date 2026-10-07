import sys
import io
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
import app
import predict
import financial_engine

def run_tests():
    print("=== 1. Testing Prediction & Sizing Logic ===")
    raw = predict.analyze_location(13.8264, 99.8780)
    assert raw['annual_kwh'] > 0, "Annual kWh should be positive"
    num_sys = len(raw['systems'])
    assert num_sys == 5, f"Expected 5 systems, got {num_sys}"

    test_cases = [
        (1500, 3),
        (3000, 3),
        (5000, 5),
        (10000, 10),
        (15000, 15),
        (20000, 20),
    ]

    for bill, expected_kw in test_cases:
        res = financial_engine.calculate_financials(raw, bill)
        rec = res['recommended']
        kw = rec['kw']
        passed = (kw == expected_kw)
        print(f"Bill: {bill:>5} THB -> Recommended: {kw:>2} kW (Expected: {expected_kw:>2} kW) | Save/mo: {rec['monthly_save']:>5} THB | ROI: {rec['roi']:>4.1f}% | Payback: {rec['breakeven']:>3.1f}y | Pass: {passed}")
        assert passed, f"Failed for bill {bill}: got {kw} kW, expected {expected_kw} kW"

    print("All solar sizing test cases passed flawlessly!")

    print("\n=== 2. Testing Flask Endpoints with test_client ===")
    with app.app.test_client() as c:
        # Public routes
        for route in ['/', '/map', '/history', '/report', '/api/stats', '/api/model-info', '/api/settings', '/api/boundary']:
            r = c.get(route)
            print(f"{route:<22} -> status {r.status_code}")
            assert r.status_code == 200, f"Route {route} failed with {r.status_code}"

        # Protected settings route check (redirects unauthenticated)
        r_set = c.get('/settings')
        print(f"{'/settings (unauth)':<22} -> status {r_set.status_code} (Protected redirect to admin login)")
        assert r_set.status_code == 302

        # Log in as admin
        with c.session_transaction() as sess:
            sess['is_admin'] = True
            sess['admin_username'] = 'admin'

        r_set_admin = c.get('/settings')
        print(f"{'/settings (admin)':<22} -> status {r_set_admin.status_code}")
        assert r_set_admin.status_code == 200

        # Test settings update with admin
        post_res = c.post('/api/settings', json={'electricity_rate': 4.50})
        assert post_res.status_code == 200

        # Test analysis endpoint
        an_res = c.post('/api/analyze', json={'lat': 13.8264, 'lng': 99.8780, 'address': 'บ้านโป่ง ทดสอบระบบ', 'monthly_bill': 5000})
        assert an_res.status_code == 200
        an_data = an_res.json
        an_id = an_data['id']
        assert an_data['recommended']['kw'] == 5

        # Test analysis page
        detail_res = c.get(f'/analysis/{an_id}')
        assert detail_res.status_code == 200
        print(f"/analysis/{an_id:<13} -> status {detail_res.status_code}")

        # Test Excel exports
        ex_single = c.get(f'/api/export/excel/{an_id}')
        assert ex_single.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(ex_single.data))
        assert 'สรุปผลการวิเคราะห์' in wb.sheetnames
        assert 'พยากรณ์รายเดือน' in wb.sheetnames
        assert 'เปรียบเทียบขนาดระบบ' in wb.sheetnames
        print(f"Excel single ({an_id}): {len(ex_single.data)} bytes, sheets: {wb.sheetnames}")

        ex_all = c.get('/api/export/excel/all')
        assert ex_all.status_code == 200
        wb_all = openpyxl.load_workbook(io.BytesIO(ex_all.data))
        print(f"Excel all: {len(ex_all.data)} bytes, sheets: {wb_all.sheetnames}")

        # Clean up test entry
        c.delete(f'/api/history/{an_id}')
        print(f"Cleaned up test record {an_id}")

    print("\n==========================================")
    print("ALL TESTS PASSED WITH 100% SUCCESS!")
    print("==========================================")

if __name__ == '__main__':
    run_tests()
